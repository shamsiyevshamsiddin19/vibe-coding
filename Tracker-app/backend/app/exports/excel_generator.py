import io
from datetime import datetime
from typing import Any, Dict, List
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def generate_excel_report(
    title: str,
    teacher_name: str,
    target_name: str,
    report_type: str,
    stats: Dict[str, Any],
    items: List[Dict[str, Any]],
) -> io.BytesIO:
    """Generate structured and professionally styled Excel report using openpyxl."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hisobot"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    title_font = Font(name="Calibri", size=16, bold=True, color="1E293B")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="64748B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True, color="1E293B")
    regular_font = Font(name="Calibri", size=10, color="1E293B")

    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    summary_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # 1. Header Title
    ws.merge_cells("A1:F1")
    ws["A1"] = f"TRACKER APP — {title.upper()}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells("A2:F2")
    ws["A2"] = f"O'qituvchi: {teacher_name} | {report_type}: {target_name} | Sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(vertical="center")

    # 2. Stats Block
    ws.append([])
    ws.append(["Ko'rsatkich", "Qiymat", "", "Ko'rsatkich", "Qiymat", ""])
    stat_row = ws.max_row
    for col in range(1, 7):
        ws.cell(row=stat_row, column=col).font = bold_font
        ws.cell(row=stat_row, column=col).fill = summary_fill
        ws.cell(row=stat_row, column=col).border = thin_border

    ws.append([
        "Jami talabalar:", f"{stats.get('total_students', 0)} ta", "",
        "O'rtacha ball:", f"{stats.get('avg_score', 0)} / 100", ""
    ])
    ws.append([
        "Bajarilgan topshiriqlar:", f"{stats.get('completed_tasks', 0)} ta", "",
        "O'zlashtirish darajasi:", f"{stats.get('completion_rate', 0)}%", ""
    ])

    for r in range(stat_row + 1, ws.max_row + 1):
        for col in range(1, 7):
            cell = ws.cell(row=r, column=col)
            cell.font = regular_font
            cell.border = thin_border

    ws.append([])  # empty spacer

    # 3. Main Data Table
    headers = ["№", "Talaba / Topshiriq", "Guruh", "Holati", "Muddat", "Baho"]
    ws.append(headers)
    header_row_idx = ws.max_row
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 4, 5, 6] else "left")
        cell.border = thin_border

    for idx, item in enumerate(items, 1):
        row_data = [
            idx,
            item.get("title", "-"),
            item.get("group", "-"),
            item.get("status", "-"),
            str(item.get("deadline", "-")),
            item.get("score", "-"),
        ]
        ws.append(row_data)
        curr_row = ws.max_row
        for col_idx in range(1, 7):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if idx % 2 == 0:
                cell.fill = alt_row_fill
            if col_idx in [1, 4, 5, 6]:
                cell.alignment = Alignment(horizontal="center")

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
